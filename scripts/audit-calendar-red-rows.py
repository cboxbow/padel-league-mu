import json
import re
import sys
from pathlib import Path
from datetime import datetime

import openpyxl


def color_text(color):
    if not color:
        return ""
    if color.type == "rgb":
        return str(color.rgb or "").upper()
    if color.type == "indexed":
        return f"indexed:{color.indexed}"
    if color.type == "theme":
        return f"theme:{color.theme}"
    return str(color.type)


def is_redish(cell):
    fill = color_text(cell.fill.fgColor)
    font = color_text(cell.font.color)
    red_values = {
        "FFFF0000",
        "FFC00000",
        "FFCC0000",
        "FFFF6666",
        "FFFF9999",
        "FFEA9999",
        "FFFFC7CE",
        "FFF4CCCC",
    }
    return fill in red_values or font in red_values


def norm(value):
    text = str(value or "").strip().upper()
    replacements = {
        "À": "A",
        "Â": "A",
        "Ä": "A",
        "Ç": "C",
        "É": "E",
        "È": "E",
        "Ê": "E",
        "Ë": "E",
        "Î": "I",
        "Ï": "I",
        "Ô": "O",
        "Ö": "O",
        "Ñ": "N",
        "Ù": "U",
        "Û": "U",
        "Ü": "U",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def club_key(value):
    text = norm(value)
    if "CANA" in text:
        return "CANA BEAU PLAN"
    if "CLUB MED" in text:
        return "CLUB MED ALBION"
    if "PORT CHAMBLY" in text:
        return "I PADEL BY RM PORT CHAMBLY"
    if "HENNESSY" in text or "HENESSY" in text:
        return "I PADEL BY RM HENNESSY"
    if "LABOURDONNAIS" in text or "LSC" in text:
        return "LABOURDONNAIS MAPOU"
    if "TERRES BRUNES" in text:
        return "TERRES BRUNES SPORTS LEISURE"
    if "SPARC" in text:
        return "SPARC CASCAVELLE"
    if "AZURI" in text:
        return "STUDIO BY RM AZURI"
    if "OXYGEN" in text:
        return "OXYGEN MOKA"
    if "SYNERGY" in text or "MOKA RANGERS" in text:
        return "MOKA RANGERS"
    if "MONT CHOISY" in text:
        return "MONT CHOISY GOLF"
    if "URBAN SPORT BLACK RIVER" in text or "URBAN BR" in text:
        return "URBAN SPORT BLACK RIVER"
    if "URBAN SPORT GRAND BAIE" in text or "URBAN GB" in text:
        return "URBAN SPORT GRAND BAIE"
    if "RM CLUB TAMARIN" in text or text == "RM T":
        return "RM CLUB TAMARIN"
    if "RM CLUB GRAND BAIE" in text or "RM GB" in text or "RM FORBACH" in text:
        return "RM CLUB GRAND BAIE"
    return text


def category_key(value):
    text = norm(value)
    if "JUNIOR" in text or re.search(r"\bU1[135]\b", text):
        return "JUNIOR"
    if "MIXED" in text or "MIXTE" in text:
        return "MIXED"
    match = re.search(r"\bM(25|50|100|250|500|1000)\b", text)
    return f"M{match.group(1)}" if match else text


def division_keys(value):
    text = norm(value)
    if "JUNIOR" in text:
        return ["junior"]
    if "MIXED" in text or "MIXTE" in text:
        return ["mixed"]
    if "MEN WOMEN" in text or "MEN AND WOMEN" in text:
        return ["men", "women"]
    if "HOMMES DAMES" in text or "HOMMES FEMMES" in text:
        return ["men", "women"]
    if "WOMEN" in text or "DAMES" in text or "FEMMES" in text:
        return ["women"]
    if "MEN" in text or "HOMMES" in text:
        return ["men"]
    return []


def date_text(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value or "")[:10]


def parse_ts_tournaments(path):
    text = path.read_text(encoding="utf-8")
    raw_blocks = re.findall(r"\{id:'([^']+)'.*?\}", text, flags=re.S)
    tournaments = []
    for block_id in raw_blocks:
        start = text.index("{", max(0, text.find(f"id:'{block_id}'") - 10))
        end = text.index("}", start)
        block = text[start:end]

        def field(name):
            match = re.search(rf"{name}:'([^']*)'", block)
            return match.group(1) if match else ""

        item = {
            "id": block_id,
            "name": field("name"),
            "date": field("date"),
            "club": field("club_name"),
            "category": field("category"),
            "division": field("division"),
            "status": field("status"),
        }
        item["clubKey"] = club_key(item["club"] or item["name"])
        item["categoryKey"] = category_key(item["category"] or item["name"])
        item["divisionKey"] = norm(item["division"]).lower()
        if item["date"] and item["division"]:
            tournaments.append(item)
    return tournaments


def matches_category(red_category, tournament):
    if red_category == "JUNIOR":
        return tournament["divisionKey"] == "junior"
    return red_category == tournament["categoryKey"]


def main():
    source = Path(sys.argv[1])
    ts_source = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb["DATABASE"]
    headers = [cell.value for cell in ws[1]]
    rows = []

    for row in ws.iter_rows(min_row=2):
        values = [cell.value for cell in row]
        if not any(values):
            continue
        red_cells = [cell.coordinate for cell in row if is_redish(cell)]
        if not red_cells:
            continue
        item = dict(zip(headers, values))
        item["row"] = row[0].row
        item["red_cells"] = red_cells
        item["dateKey"] = date_text(item.get("DATE"))
        item["clubKey"] = club_key(item.get("CLUB"))
        item["categoryKey"] = category_key(item.get("CATEGORY") or item.get("CATEGORIE"))
        item["divisionKeys"] = division_keys(item.get("TYPE"))
        rows.append(item)

    output = {"redRows": len(rows), "rows": rows}

    if ts_source:
        tournaments = parse_ts_tournaments(ts_source)
        matched = []
        unmatched = []
        for row in rows:
            wanted_divisions = row["divisionKeys"]
            exact = [
                tournament for tournament in tournaments
                if tournament["date"] == row["dateKey"]
                and tournament["clubKey"] == row["clubKey"]
                and matches_category(row["categoryKey"], tournament)
                and (not wanted_divisions or tournament["divisionKey"] in wanted_divisions)
            ]
            if not exact:
                red_date = datetime.strptime(row["dateKey"], "%Y-%m-%d")
                exact = [
                    tournament for tournament in tournaments
                    if abs((datetime.strptime(tournament["date"], "%Y-%m-%d") - red_date).days) <= 1
                    and tournament["clubKey"] == row["clubKey"]
                    and matches_category(row["categoryKey"], tournament)
                    and (not wanted_divisions or tournament["divisionKey"] in wanted_divisions)
                ]
            if not exact:
                fuzzy = [
                    tournament for tournament in tournaments
                    if tournament["clubKey"] == row["clubKey"]
                    and matches_category(row["categoryKey"], tournament)
                    and (not wanted_divisions or tournament["divisionKey"] in wanted_divisions)
                ]
                row["fuzzyMatches"] = fuzzy
                unmatched.append(row)
                continue
            matched.extend(exact)

        rules = {}
        for tournament in matched:
            key = (tournament["date"], tournament["club"], tournament["category"])
            rule = rules.setdefault(
                key,
                {"date": tournament["date"], "club": tournament["club"], "category": tournament["category"], "divisions": []},
            )
            if tournament["divisionKey"] not in rule["divisions"]:
                rule["divisions"].append(tournament["divisionKey"])

        output["matchedTournamentCount"] = len(matched)
        output["matchedTournamentIds"] = [tournament["id"] for tournament in matched]
        output["unmatchedRows"] = unmatched
        output["rules"] = list(rules.values())

    print(json.dumps(output, default=str, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
